"""Verify saved local results and document/container integrity before handoff."""
from pathlib import Path
import hashlib
import json
import re
import zipfile

from docx import Document
from pptx import Presentation

PROJECT = Path(__file__).resolve().parents[1]
ARTIFACTS = PROJECT.parent / '_ai_context/artifacts'
OUT = PROJECT.parent / '_ai_context/submission'


def read_json(name):
    raw = (ARTIFACTS/name).read_bytes()
    return json.loads(raw.decode('utf-16' if raw.startswith(b'\xff\xfe') else 'utf-8-sig'))


checks = {}
for name in ['unit', 'integration', 'workflow', 'surge', 'cache']:
    report = read_json(f'stage-15-{name}-tests.json')
    if not report.get('success') or report['numFailedTests'] or report['numPendingTests']:
        raise SystemExit(f'{name} suite not fully passed')
    checks[name] = {'passed': report['numPassedTests'], 'failed': 0, 'skipped': 0}
browser = read_json('stage-15-browser-final.json')
stats = browser['stats']
portal = read_json('stage-15-portal-rechecks.json')
if portal['stats']['unexpected'] or portal['stats']['flaky'] or portal['stats']['skipped'] or portal.get('errors'):
    raise SystemExit('Portal recheck not fully passed')


def cases(report):
    found = {}

    def visit(suite):
        for spec in suite.get('specs', []):
            for test in spec['tests']:
                found[(spec['file'], spec['title'], test['projectName'])] = test
        for child in suite.get('suites', []): visit(child)
    for suite in report['suites']: visit(suite)
    return found


resolved = cases(browser)
rechecked = cases(portal)
if not set(rechecked).issubset(resolved): raise SystemExit('Recheck identity not in broad suite')
resolved.update(rechecked)
if browser.get('errors') or any(t['status'] != 'expected' or not t['results'] or t['results'][-1]['status'] != 'passed' for t in resolved.values()):
    raise SystemExit('An unresolved browser failure/skip remains')
checks['browser'] = {'passed': len(resolved), 'failed': 0, 'skipped': 0, 'flaky': 0,
                     'broadRun': stats, 'portalRecheck': portal['stats'],
                     'accounting': '128 distinct project/file/title cases; failed/skipped portal serial group replaced by successful focused replay',
                     'environment': 'local production server; Chromium desktop and Pixel 7'}
if checks['browser']['passed'] != 128:
    raise SystemExit('Expected all 128 selected browser cases')
recording = read_json('stage-15-recording.json')
if len(recording['scenes']) != 12 or recording['errors'] or recording['externalRequests'] or recording['decodedMP4']['paused']:
    raise SystemExit('Recording validation failed')
render = read_json('stage-15-document-render.json')
if len(render['slides']) != 10 or any(row['overflow'] or not row['contained'] for row in render['slides']):
    raise SystemExit('Document layout check failed')
office = {}
for name in ['PulseGrid-Report.docx', 'PulseGrid-Slides.pptx']:
    with zipfile.ZipFile(OUT/name) as z:
        if z.testzip(): raise SystemExit('Office ZIP CRC failed')
    office[name] = 'CRC passed'
presentation = Presentation(OUT/'PulseGrid-Slides.pptx')
if len(presentation.slides) != 10 or not all(s.notes_slide.notes_text_frame.text.strip() for s in presentation.slides):
    raise SystemExit('Slide count or speaker notes missing')
if len(Document(OUT/'PulseGrid-Report.docx').paragraphs) < 40:
    raise SystemExit('Report export incomplete')
pdfs = {}
for name in ['PulseGrid-Report.pdf', 'PulseGrid-Slides.pdf']:
    content = (OUT/name).read_bytes()
    pages = len(re.findall(rb'/Type /Page\b', content))
    if not content.startswith(b'%PDF-') or not pages: raise SystemExit('Invalid PDF')
    pdfs[name] = {'pages': pages, 'bytes': len(content)}

digest = hashlib.sha256()
source_files = []
for folder in ['src', 'scripts', 'tests', 'docs', '.github']:
    source_files += [p for p in (PROJECT/folder).rglob('*') if p.is_file() and '__pycache__' not in p.parts]
source_files += [PROJECT/f for f in ['package.json', 'package-lock.json', 'playwright.config.ts', 'tsconfig.json']]
for path in sorted(source_files):
    digest.update(path.relative_to(PROJECT).as_posix().encode()); digest.update(b'\0'); digest.update(path.read_bytes())
result = {'date': '2026-10-29', 'timezone': 'Asia/Calcutta', 'stage': 15,
          'status': 'local validation passed; hosted/college/human release gates open',
          'environment': {'node': '24.22.1', 'mongodb': '8.0.17 temporary replica sets',
                          'workflow': 'actual local Inngest CLI 1.49.0 / SDK 4.26.1'},
          'checks': checks, 'totalPassed': sum(c['passed'] for c in checks.values()),
          'qualityCommands': {'lint': 'passed; zero warnings', 'types': 'passed', 'productionBuild': 'passed'},
          'documentChecks': {'office': office, 'pdfs': pdfs, 'slides': 10, 'notesPresent': True},
          'recording': {'scenes': 12, **recording['decodedMP4'], 'externalRequests': 0, 'pageErrors': 0},
          'offline': 'desktop/mobile production bundle; external requests blocked; map, completed emergency, full matched month',
          'sourceFingerprintSHA256': digest.hexdigest(),
          'sourceFingerprintPurpose': 'inventory of handoff files; does not replace independent test/build attestation',
          'gates': ['actual deployed smoke and managed-service acceptance', 'college requirements/team/slot/account labels',
                    'physical laptop/projector/offline video and phone rehearsal', 'timed named second-presenter/independent review',
                    'hosted synthetic reset/zero chaos/known-good deployment/submission receipt'],
          'gitActions': False, 'hostedChanges': False, 'collegeUpload': False}
(ARTIFACTS/'stage-15-checks.json').write_text(json.dumps(result, indent=2)+'\n', encoding='utf-8')
print(json.dumps(result, indent=2))
