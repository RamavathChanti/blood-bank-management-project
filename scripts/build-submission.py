"""Export reviewable local college documents from the Markdown sources; no network."""
from pathlib import Path
import base64
import html
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / '_ai_context' / 'submission'
OUT.mkdir(parents=True, exist_ok=True)
DISCLAIMER = 'Teaching model, not clinical software.'


def plain(text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text.replace('**', '').replace('`', '')


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            yield 'code', '\n'.join(code)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s:|\-]+$', lines[i]):
                    rows.append([s.strip() for s in lines[i].strip('|').split('|')])
                i += 1
            yield 'table', rows
            continue
        elif re.match(r'^#{1,6} ', line):
            level = len(line.split(' ')[0])
            yield f'h{level}', line[level + 1:]
        elif line.startswith('- '):
            yield 'li', line[2:]
        else:
            para = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|- )', lines[i]):
                para.append(lines[i]); i += 1
            yield 'p', ' '.join(para)
            continue
        i += 1


def rich_html(text):
    text = html.escape(text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'<a href="\2">\1</a>', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'<strong>\1</strong>', text)
    return re.sub(r'`([^`]+)`', r'<code>\1</code>', text)


def doc_html(text):
    parts = []
    for kind, value in blocks(text):
        if kind == 'code': parts.append('<pre>' + html.escape(value) + '</pre>')
        elif kind == 'table':
            rows = ['<tr>' + ''.join(f'<{"th" if i == 0 else "td"}>{rich_html(cell)}</{"th" if i == 0 else "td"}>' for cell in row) + '</tr>' for i, row in enumerate(value)]
            parts.append('<table>' + ''.join(rows) + '</table>')
        elif kind == 'li': parts.append('<p class="bullet">• ' + rich_html(value) + '</p>')
        else: parts.append(f'<{kind}>' + rich_html(value) + f'</{kind}>')
    return '\n'.join(parts)


def diagram():
    # One logical diagram, exported as vector and high-resolution bitmap for Office.
    boxes = [
        (20, 90, 350, 100, ['Private portal browsers', 'Bank / hospital / donor', 'State-admin session']),
        (450, 90, 350, 100, ['Clerk identity (IDaaS)', 'Verified current membership']),
        (880, 90, 350, 100, ['Separate browser workers', 'Seeded clock / offline labs', 'No Atlas access']),
        (450, 260, 350, 110, ['Next.js / Vercel Node API', 'Validation / authorization', 'Tenant scope / traces']),
        (20, 450, 350, 100, ['Atlas registry (DBaaS)', 'Transactions / ownership', 'Counters / signed audit']),
        (450, 450, 350, 100, ['Inngest workflows', 'Durable steps / retries', 'Acknowledgements / waits']),
        (880, 450, 350, 100, ['Public Board bands', 'Allowlisted aggregate', 'Next cache / intended CDN']),
    ]
    edges = [(370, 135, 450, 135), (195, 190, 450, 300), (620, 190, 620, 260), (450, 340, 195, 450), (620, 370, 620, 450), (800, 340, 1055, 450)]
    svg = ['<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="600" viewBox="0 0 1280 600"><defs><marker id="arrow" markerWidth="10" markerHeight="8" refX="9" refY="4" orient="auto"><path d="M0 0L10 4L0 8" fill="#75d7bb"/></marker></defs><rect width="1280" height="600" fill="#0c1825"/><g font-family="Arial,sans-serif" fill="#f0f5fa"><text x="40" y="40" font-size="26" font-weight="bold">PulseGrid — implementation architecture</text>']
    image = Image.new('RGB', (2560, 1200), '#0c1825'); draw = ImageDraw.Draw(image)
    font_path = Path('C:/Windows/Fonts/arial.ttf')
    font = ImageFont.truetype(str(font_path), 46) if font_path.exists() else ImageFont.load_default()
    large = ImageFont.truetype(str(font_path), 48) if font_path.exists() else font
    draw.text((80, 24), 'PulseGrid — implementation architecture', fill='#f0f5fa', font=large)
    for x1, y1, x2, y2 in edges:
        svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#75d7bb" stroke-width="3" marker-end="url(#arrow)"/>')
        draw.line((x1*2, y1*2, x2*2, y2*2), fill='#75d7bb', width=6)
    for x, y, w, h, labels in boxes:
        svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="12" fill="#182b3c" stroke="#568a93"/>')
        draw.rounded_rectangle((x*2, y*2, (x+w)*2, (y+h)*2), radius=24, fill='#182b3c', outline='#568a93', width=3)
        for j, label in enumerate(labels):
            sy = y + 29 + j*28
            svg.append(f'<text x="{x+16}" y="{sy}" font-size="23">{escape(label)}</text>')
            draw.text(((x+16)*2, (sy-18)*2), label, fill='#f0f5fa', font=font)
    foot = ['Teaching model, not clinical software. Implemented code; hosted acceptance pending.']
    for i, label in enumerate(foot):
        svg.append(f'<text x="20" y="590" font-size="22">{escape(label)}</text>')
        draw.text((40, 1136), label, fill='#d4e5ec', font=font)
    svg.append('</g></svg>')
    (OUT/'architecture.svg').write_text(''.join(svg), encoding='utf-8')
    image.save(OUT/'architecture.png')


def make_report():
    text = (OUT/'REPORT.md').read_text(encoding='utf-8')
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    section.header.paragraphs[0].text = 'PULSEGRID  /  COLLEGE DEMO  /  LOCAL EVIDENCE CANDIDATE'
    footer = section.footer.paragraphs[0]; footer.text = DISCLAIMER + '  |  Page '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for kind, value in blocks(text):
        if kind.startswith('h'):
            doc.add_heading(plain(value), min(int(kind[1:]), 3))
        elif kind == 'table':
            table = doc.add_table(rows=0, cols=len(value[0])); table.style = 'Light Shading Accent 1'
            for i, row in enumerate(value):
                cells = table.add_row().cells
                for j, cell in enumerate(row): cells[j].text = plain(cell)
                if i == 0:
                    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        else:
            p = doc.add_paragraph(plain(value), 'List Bullet' if kind == 'li' else None)
            if kind == 'code':
                for run in p.runs: run.font.name = 'Consolas'; run.font.size = Pt(8)
    doc.add_page_break(); doc.add_heading('Architecture', 1); doc.add_picture(str(OUT/'architecture.png'), width=Inches(6.6))
    doc.save(OUT/'PulseGrid-Report.docx')
    css = '''body{font:11pt/1.5 Arial;color:#142a3b;margin:0}h1{font-size:29pt;color:#12394d}h2{font-size:17pt;color:#126b70;margin-top:24pt;break-after:avoid}h3{break-after:avoid}p{orphans:3;widows:3}table{border-collapse:collapse;width:100%;font-size:9pt;margin:14pt 0}td,th{border:1px solid #b7cbd2;padding:6pt;text-align:left;overflow-wrap:anywhere}th{background:#e5f0f0}tr{break-inside:avoid}a{color:#125e9a;overflow-wrap:anywhere}pre{white-space:pre-wrap;background:#eef3f5;padding:10pt}.diagram{page-break-before:always}.diagram img{width:100%}'''
    encoded = base64.b64encode((OUT/'architecture.svg').read_bytes()).decode()
    page = '<!doctype html><html lang="en"><meta charset="utf-8"><title>PulseGrid Report</title><style>'+css+'</style><body>'+doc_html(text)+f'<section class="diagram"><h2>Architecture</h2><img alt="PulseGrid architecture" src="data:image/svg+xml;base64,{encoded}"></section></body></html>'
    (OUT/'report.html').write_text(page, encoding='utf-8')


def slide_data():
    text = (OUT/'SLIDES.md').read_text(encoding='utf-8')
    slides = []
    for match in re.finditer(r'^## ([^\n]+)\n(.*?)(?=^## |\Z)', text, re.M | re.S):
        title, content = match.groups()
        bullets = re.findall(r'^- (.+)', content, re.M)
        notes = re.search(r'Notes: (.+)', content, re.S)
        slides.append((title, bullets, notes.group(1).strip() if notes else ''))
    return slides


def make_slides():
    prs = Presentation(); prs.slide_width = SlideInches(13.333); prs.slide_height = SlideInches(7.5)
    html_slides = []
    encoded = base64.b64encode((OUT/'architecture.svg').read_bytes()).decode()
    for idx, (title, bullets, notes) in enumerate(slide_data()):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string('0C1825')

        def text_box(x, y, w, h, content, size, color='F0F5FA', bold=False):
            box = slide.shapes.add_textbox(SlideInches(x), SlideInches(y), SlideInches(w), SlideInches(h))
            tf = box.text_frame; tf.word_wrap = True
            tf.margin_left = tf.margin_right = 0
            for i, line in enumerate(content):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = line
                p.font.name = 'Arial'; p.font.size = SlidePt(size); p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color); p.space_after = SlidePt(20)

        text_box(.6, .25, 11, .4, ['PULSEGRID / COLLEGE DEMO / SYNTHETIC ONLY'], 13, '75D7BB')
        text_box(.6, .85, 12, 1.2, [title], 32, bold=True)
        if idx == 2:
            slide.shapes.add_picture(str(OUT/'architecture.png'), SlideInches(.7), SlideInches(1.7), width=SlideInches(10.8))
            # The diagram occupies the slide; explanatory bullets remain in speaker notes.
        else:
            text_box(.8, 2.4, 11.7, 3.5, ['• '+line for line in bullets], 28)
        text_box(.6, 6.95, 11.7, .35, [DISCLAIMER+'  |  Selected architecture; hosted acceptance pending.'], 12, 'C1D3DF')
        text_box(12.4, 6.95, .5, .35, [str(idx+1)], 12, '75D7BB')
        slide.notes_slide.notes_text_frame.text = ('\n'.join(bullets)+'\n\n'+notes)
        diagram_html = f'<img src="data:image/svg+xml;base64,{encoded}" alt="Architecture">' if idx == 2 else ''
        html_slides.append(f'<section class="slide"><div class="eyebrow">PULSEGRID / COLLEGE DEMO / SYNTHETIC ONLY</div><h1>{html.escape(title)}</h1><div class="{"architecture" if idx == 2 else "body"}">{diagram_html}<ul>'+''.join('<li>'+html.escape(s)+'</li>' for s in ([] if idx == 2 else bullets))+f'</ul></div><footer>{DISCLAIMER} | Selected architecture; hosted acceptance pending.<span>{idx+1}</span></footer></section>')
    prs.save(OUT/'PulseGrid-Slides.pptx')
    css = '''@page{size:13.333in 7.5in;margin:0}*{box-sizing:border-box}body{margin:0;background:#0c1825;color:#f0f5fa;font-family:Arial}.slide{width:1280px;height:720px;padding:42px 58px;position:relative;page-break-after:always;overflow:hidden}.slide:last-child{page-break-after:auto}.eyebrow{font-size:15px;letter-spacing:2px;color:#75d7bb}h1{font-size:39px;line-height:1.2;margin:40px 0 34px}li{font-size:30px;line-height:1.45;margin-bottom:28px}ul{padding-left:30px;margin-top:0}.architecture{display:flex;align-items:center;justify-content:center;margin-top:-22px}.architecture img{width:1050px}.architecture ul{width:450px}.architecture li{font-size:23px;margin-bottom:22px}footer{position:absolute;bottom:30px;left:58px;right:58px;color:#c1d3df;font-size:14px}footer span{float:right}'''
    (OUT/'slides.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><title>PulseGrid Slides</title><style>'+css+'</style><body>'+''.join(html_slides)+'</body></html>', encoding='utf-8')


diagram()
make_report()
make_slides()
print('Generated architecture.svg/png, report DOCX/HTML, and 10-slide PPTX/HTML with notes. Render PDFs using node scripts/render-submission.mjs.')
